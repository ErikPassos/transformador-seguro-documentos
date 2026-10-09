from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

def ajustar_larguras_colunas(planilha):
    """
    Ajusta automaticamente a largura das colunas
    de uma planilha do Excel.
    """

    for coluna in planilha.columns:
        maior_tamanho = 0

        numero_coluna = coluna[0].column

        letra_coluna = get_column_letter(
            numero_coluna
        )

        for celula in coluna:
            if celula.value is None:
                continue

            tamanho_conteudo = len(
                str(celula.value)
            )

            if tamanho_conteudo > maior_tamanho:
                maior_tamanho = tamanho_conteudo

        largura_ajustada = min(
            maior_tamanho + 3,
            50,
        )

        planilha.column_dimensions[
            letra_coluna
        ].width = largura_ajustada


def formatar_planilha(planilha):
    """
    Aplica formatação básica à planilha.
    """

    cor_cabecalho = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    fonte_cabecalho = Font(
        color="FFFFFF",
        bold=True,
    )

    for celula in planilha[1]:
        celula.fill = cor_cabecalho
        celula.font = fonte_cabecalho
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = (
        planilha.dimensions
    )

    for linha in planilha.iter_rows():
        for celula in linha:
            celula.alignment = Alignment(
                vertical="center",
            )

    ajustar_larguras_colunas(planilha)


def criar_caminho_saida(
    caminho_pdf,
    pasta_saida,
):
    """
    Cria um nome único para o arquivo Excel.

    O nome utiliza o arquivo PDF original e
    a data e hora da exportação.
    """

    caminho_pdf = Path(caminho_pdf)
    pasta_saida = Path(pasta_saida)

    pasta_saida.mkdir(
        parents=True,
        exist_ok=True,
    )

    momento_exportacao = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    nome_base = caminho_pdf.stem

    nome_excel = (
        f"{nome_base}_extraido_"
        f"{momento_exportacao}.xlsx"
    )

    return pasta_saida / nome_excel


def exportar_tabelas_excel(
    tabelas_extraidas,
    caminho_pdf,
    pasta_saida="saida",
):
    """
    Exporta as tabelas extraídas para um arquivo Excel.

    Cada tabela é salva em uma planilha diferente.

    Retorna:
        tuple:
            bool:
                indica se a exportação foi concluída;

            Path ou None:
                caminho do arquivo criado;

            str:
                mensagem explicativa.
    """

    if not isinstance(tabelas_extraidas, list):
        return (
            False,
            None,
            "As tabelas devem ser informadas "
            "em uma lista.",
        )

    if len(tabelas_extraidas) == 0:
        return (
            False,
            None,
            "Não existem tabelas para exportar.",
        )

    try:
        caminho_saida = criar_caminho_saida(
            caminho_pdf,
            pasta_saida,
        )

        with __import__("pandas").ExcelWriter(
            caminho_saida,
            engine="openpyxl",
        ) as escritor_excel:

            nomes_planilhas = set()

            for tabela in tabelas_extraidas:
                dataframe = tabela.get("dataframe")

                if dataframe is None:
                    continue

                pagina = tabela.get(
                    "pagina",
                    0,
                )

                numero_tabela = tabela.get(
                    "numero_tabela",
                    0,
                )

                nome_planilha = (
                    f"Pag{pagina}_Tabela"
                    f"{numero_tabela}"
                )

                nome_planilha = nome_planilha[:31]

                nome_original = nome_planilha
                contador = 2

                while nome_planilha in nomes_planilhas:
                    sufixo = f"_{contador}"

                    nome_planilha = (
                        nome_original[
                            : 31 - len(sufixo)
                        ]
                        + sufixo
                    )

                    contador += 1

                nomes_planilhas.add(
                    nome_planilha
                )

                dataframe.to_excel(
                    escritor_excel,
                    sheet_name=nome_planilha,
                    index=False,
                )

                planilha = escritor_excel.sheets[
                    nome_planilha
                ]

                formatar_planilha(planilha)

        return (
            True,
            caminho_saida,
            "Exportação para Excel concluída "
            "com sucesso.",
        )

    except PermissionError:
        return (
            False,
            None,
            "O programa não possui permissão "
            "para criar o arquivo Excel. "
            "Verifique se o arquivo está aberto.",
        )

    except OSError as erro:
        return (
            False,
            None,
            "O sistema operacional não conseguiu "
            "criar o arquivo Excel. "
            f"Detalhe técnico: {erro}",
        )

    except Exception as erro:
        return (
            False,
            None,
            "Não foi possível exportar as tabelas. "
            f"Tipo do erro: {type(erro).__name__}. "
            f"Detalhe técnico: {erro}",
        )